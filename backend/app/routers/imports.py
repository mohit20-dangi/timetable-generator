from datetime import datetime, time
import csv
import io
import zipfile
from typing import Any, Dict, Iterable, List, Mapping, Type

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from pydantic import ValidationError

from app.auth.dependencies import require_admin
from app.db import get_db
from app.models import (
    AcademicYear,
    Department,
    ElectiveGroup,
    LabBatch,
    Prerequisite,
    Room,
    Section,
    SectionSubject,
    Subject,
    SubjectType,
    Teacher,
    TeacherSubject,
    TimeSlot,
    User,
)
from app.schemas import (
    AcademicYearCreate,
    ElectiveGroupCreate,
    LabBatchCreate,
    PrerequisiteCreate,
    RoomCreate,
    SectionCreate,
    SectionSubjectCreate,
    SubjectCreate,
    TeacherCreate,
    TeacherSubjectCreate,
    TimeSlotCreate,
)
from app.schemas.room import VALID_ROOM_TYPES


router = APIRouter(prefix="/api/import", tags=["Bulk Import"])


SHEET_DEFINITIONS: Dict[str, Dict[str, Any]] = {
    "AcademicYears": {
        "model": AcademicYear,
        "required": {"id", "name"},
        "columns": ["id", "name", "department_id", "num_sections", "default_section_strength"],
        "int": {"num_sections", "default_section_strength"},
        "schema": AcademicYearCreate,
    },
    "Sections": {
        "model": Section,
        "required": {"id", "year_id", "name"},
        "columns": ["id", "year_id", "name", "strength"],
        "int": {"strength"},
        "schema": SectionCreate,
    },
    "ElectiveGroups": {
        "model": ElectiveGroup,
        "required": {"id", "department_id", "name"},
        "columns": ["id", "department_id", "name", "must_be_parallel"],
        "bool": {"must_be_parallel"},
        "schema": ElectiveGroupCreate,
    },
    "Subjects": {
        "model": Subject,
        "required": {"id", "name", "type"},
        "columns": [
            "id", "name", "code", "type", "department_id", "category", "delivery_mode",
            "lecture_hours", "tutorial_hours", "practical_hours",
            "scheme_hours_per_week", "weekly_hours", "sessions_per_week",
            "periods_per_session", "back_to_back", "batch_scheduling_mode",
            "max_per_day", "requires_room_type",
            "linked_group_id", "elective_group_id",
        ],
        "int": {
            "lecture_hours", "tutorial_hours", "practical_hours", "weekly_hours",
            "sessions_per_week", "periods_per_session", "max_per_day", "scheme_hours_per_week",
        },
        "bool": {"back_to_back"},
        "schema": SubjectCreate,
    },
    "Teachers": {
        "model": Teacher,
        "required": {"id", "name"},
        "columns": [
            "id", "name", "initials", "department_id", "max_continuous_classes", "max_daily_classes",
            "max_weekly_hours", "is_guest_from_other_dept",
        ],
        "int": {"max_continuous_classes", "max_daily_classes", "max_weekly_hours"},
        "bool": {"is_guest_from_other_dept"},
        "schema": TeacherCreate,
    },
    "Rooms": {
        "model": Room,
        "required": {"id", "name", "type"},
        "columns": ["id", "name", "type", "capacity", "department_id"],
        "int": {"capacity"},
        "schema": RoomCreate,
    },
    "SectionSubjects": {
        "model": SectionSubject,
        "required": {"section_id", "subject_id"},
        "columns": ["section_id", "subject_id", "is_elective", "elective_group_id"],
        "bool": {"is_elective"},
        "schema": SectionSubjectCreate,
    },
    "TeacherSubjects": {
        "model": TeacherSubject,
        "required": {"teacher_id", "subject_id"},
        "columns": ["teacher_id", "subject_id"],
        "schema": TeacherSubjectCreate,
    },
    "LabBatches": {
        "model": LabBatch,
        "required": {"id", "section_id", "batch_name"},
        "columns": ["id", "section_id", "batch_name", "strength"],
        "int": {"strength"},
        "schema": LabBatchCreate,
    },
    "Prerequisites": {
        "model": Prerequisite,
        "required": {"subject_id", "requires_subject_id"},
        "columns": ["subject_id", "requires_subject_id"],
        "schema": PrerequisiteCreate,
    },
    "TimeSlots": {
        "model": TimeSlot,
        "required": {"id", "day", "period_index", "start_time", "end_time"},
        "columns": ["id", "day", "period_index", "start_time", "end_time"],
        "int": {"period_index"},
        "time": {"start_time", "end_time"},
        "schema": TimeSlotCreate,
    },
}

# Sheets that fill in one field of a SHEET_DEFINITIONS entity, one row per
# value, instead of asking the admin to type a JSON blob into a single Excel
# cell (e.g. `{"mon": [...], "tue": [...]}`). Mirrors the existing
# SectionSubjects/TeacherSubjects pattern of modelling a many-valued
# relationship as its own sheet rather than a nested structure in one cell.
CHILD_SHEET_DEFINITIONS: Dict[str, Dict[str, Any]] = {
    "TeacherAvailability": {
        "parent_sheet": "Teachers",
        "parent_field": "availability",
        "parent_key": "teacher_id",
        "columns": ["teacher_id", "day", "start_time", "end_time"],
        "time": {"start_time", "end_time"},
        "build": lambda rows: [
            {"day": r["day"], "start": r["start_time"], "end": r["end_time"]} for r in rows
        ],
    },
    "TeacherPreferredSlots": {
        "parent_sheet": "Teachers",
        "parent_field": "preferred_slots",
        "parent_key": "teacher_id",
        "columns": ["teacher_id", "day", "start_time", "end_time"],
        "time": {"start_time", "end_time"},
        "build": lambda rows: _group_slots_by_day(rows),
    },
    "RoomAvailability": {
        "parent_sheet": "Rooms",
        "parent_field": "availability",
        "parent_key": "room_id",
        "columns": ["room_id", "day", "start_time", "end_time"],
        "time": {"start_time", "end_time"},
        "build": lambda rows: [
            {"day": r["day"], "start": r["start_time"], "end": r["end_time"]} for r in rows
        ],
    },
    "RoomEquipment": {
        "parent_sheet": "Rooms",
        "parent_field": "equipment",
        "parent_key": "room_id",
        "columns": ["room_id", "equipment"],
        "build": lambda rows: [r["equipment"] for r in rows],
    },
    "RoomSharedDepartments": {
        "parent_sheet": "Rooms",
        "parent_field": "shared_with_departments",
        "parent_key": "room_id",
        "columns": ["room_id", "department_id"],
        "build": lambda rows: [r["department_id"] for r in rows],
    },
    "SubjectEquipment": {
        "parent_sheet": "Subjects",
        "parent_field": "requires_equipment",
        "parent_key": "subject_id",
        "columns": ["subject_id", "equipment"],
        "build": lambda rows: [r["equipment"] for r in rows],
    },
    "ElectiveGroupSubjects": {
        "parent_sheet": "ElectiveGroups",
        "parent_field": "offered_subject_ids",
        "parent_key": "elective_group_id",
        "columns": ["elective_group_id", "subject_id"],
        "build": lambda rows: [r["subject_id"] for r in rows],
    },
    "AcademicYearLunchWindows": {
        "parent_sheet": "AcademicYears",
        "parent_field": "lunch_windows",
        "parent_key": "year_id",
        "columns": ["year_id", "day", "start_time", "end_time"],
        "time": {"start_time", "end_time"},
        "build": lambda rows: _build_lunch_windows(rows),
    },
}

# Secondary reference columns inside a child sheet's built value that point
# at another entity (beyond the parent_key itself), so a typo still fails
# loudly instead of silently producing an empty relationship at solve time.
CHILD_VALUE_REFERENCES = {
    "ElectiveGroupSubjects": (Subject, "subject"),
    "RoomSharedDepartments": (Department, "department"),
}


def _group_slots_by_day(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_day: Dict[str, List[Dict[str, str]]] = {}
    for row in rows:
        by_day.setdefault(row["day"], []).append({"start": row["start_time"], "end": row["end_time"]})
    return [{"day": day, "slots": slots} for day, slots in by_day.items()]


def _build_lunch_windows(rows: List[Dict[str, Any]]) -> Dict[str, List[str]]:
    windows: Dict[str, List[str]] = {}
    for row in rows:
        if row["day"] in windows:
            raise ValueError(f"AcademicYearLunchWindows: duplicate lunch window for day '{row['day']}'")
        windows[row["day"]] = [row["start_time"], row["end_time"]]
    return windows

RELATION_KEYS = {
    "SectionSubjects": (SectionSubject, ("section_id", "subject_id")),
    "TeacherSubjects": (TeacherSubject, ("teacher_id", "subject_id")),
    "Prerequisites": (Prerequisite, ("subject_id", "requires_subject_id")),
}


def _normalise_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def _cell_value(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
    return value


def _parse_time(value: Any, sheet: str, row_number: int, column: str) -> time | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    try:
        parsed = datetime.strptime(str(value).strip(), "%H:%M")
        return parsed.time()
    except ValueError as exc:
        raise ValueError(f"{sheet} row {row_number}: {column} must be HH:MM") from exc


def _parse_bool(value: Any, sheet: str, row_number: int, column: str) -> bool:
    if isinstance(value, bool):
        return value
    normalised = str(value or "").strip().lower()
    if normalised in {"true", "1", "yes", "y"}:
        return True
    if normalised in {"false", "0", "no", "n", ""}:
        return False
    raise ValueError(f"{sheet} row {row_number}: {column} must be true or false")


def _parse_int(value: Any, sheet: str, row_number: int, column: str) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{sheet} row {row_number}: {column} must be an integer") from exc


def _read_child_sheet(workbook: Any, sheet_name: str, definition: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Every column on a child sheet is required - there's no optional data
    on a row whose only purpose is "teacher X is free on Mon 9-11"."""
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_normalise_header(value) for value in rows[0]]
    present = {header for header in headers if header}
    missing = set(definition["columns"]) - present
    if missing:
        raise ValueError(f"{sheet_name}: missing required columns: {', '.join(sorted(missing))}")
    if len(headers) != len(set(headers)):
        raise ValueError(f"{sheet_name}: duplicate column headers are not allowed")

    result: List[Dict[str, Any]] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in values):
            continue
        raw = {headers[index]: _cell_value(values[index]) for index in range(len(headers)) if headers[index]}
        item: Dict[str, Any] = {}
        for column in definition["columns"]:
            value = raw.get(column)
            if column in definition.get("time", set()):
                parsed_time = _parse_time(value, sheet_name, row_number, column)
                if parsed_time is None:
                    raise ValueError(f"{sheet_name} row {row_number}: {column} is required")
                value = parsed_time.strftime("%H:%M")
            if value in (None, ""):
                raise ValueError(f"{sheet_name} row {row_number}: {column} is required")
            if column == "day" or column.endswith("_id"):
                value = str(value).strip()
            item[column] = value
        result.append(item)
    return result


def _read_child_sheets(workbook: Any) -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Returns {parent_sheet: {parent_id: {parent_field: built_value}}} by
    reading every present child sheet and grouping its rows by parent id."""
    merged: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for child_name, definition in CHILD_SHEET_DEFINITIONS.items():
        if child_name not in workbook.sheetnames:
            continue
        child_rows = _read_child_sheet(workbook, child_name, definition)
        grouped: Dict[str, List[Dict[str, Any]]] = {}
        for row in child_rows:
            grouped.setdefault(row[definition["parent_key"]], []).append(row)
        parent_sheet = definition["parent_sheet"]
        parent_field = definition["parent_field"]
        bucket = merged.setdefault(parent_sheet, {})
        for parent_id, group_rows in grouped.items():
            bucket.setdefault(parent_id, {})[parent_field] = definition["build"](group_rows)
    return merged


def _read_rows(
    workbook: Any, child_values: Dict[str, Dict[str, Dict[str, Any]]]
) -> Dict[str, List[Dict[str, Any]]]:
    parsed: Dict[str, List[Dict[str, Any]]] = {}
    for sheet_name, definition in SHEET_DEFINITIONS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_normalise_header(value) for value in rows[0]]
        present = {header for header in headers if header}
        missing = definition["required"] - present
        if missing:
            raise ValueError(f"{sheet_name}: missing required columns: {', '.join(sorted(missing))}")
        if len(headers) != len(set(headers)):
            raise ValueError(f"{sheet_name}: duplicate column headers are not allowed")

        sheet_rows: List[Dict[str, Any]] = []
        for row_number, values in enumerate(rows[1:], start=2):
            if not any(value not in (None, "") for value in values):
                continue
            raw = {headers[index]: _cell_value(values[index]) for index in range(len(headers)) if headers[index]}
            item: Dict[str, Any] = {}
            for column in definition["columns"]:
                if column not in raw:
                    continue
                value = raw[column]
                if column in definition.get("time", set()):
                    value = _parse_time(value, sheet_name, row_number, column)
                elif column in definition.get("bool", set()):
                    value = _parse_bool(value, sheet_name, row_number, column)
                elif column in definition.get("int", set()):
                    value = _parse_int(value, sheet_name, row_number, column)
                    if value is None and column not in definition["required"]:
                        # A blank optional numeric cell means "use the
                        # schema/model default" (e.g. Room.capacity=60),
                        # not "explicitly set to zero/null" - passing None
                        # here would fail Pydantic validation for fields
                        # like `int = 60` that aren't Optional.
                        continue
                if column == "id" or column.endswith("_id"):
                    value = str(value).strip() if value not in (None, "") else value
                item[column] = value
            for required in definition["required"]:
                if item.get(required) in (None, ""):
                    raise ValueError(f"{sheet_name} row {row_number}: {required} is required")
            # Pull in anything a child sheet (e.g. TeacherAvailability) built
            # for this row's id, so the rest of the pipeline sees a normal
            # availability/equipment/etc. field like it always has.
            item_id = str(item["id"]) if item.get("id") not in (None, "") else None
            bucket = child_values.get(sheet_name)
            if item_id and bucket and item_id in bucket:
                item.update(bucket.pop(item_id))
            sheet_rows.append(item)
        key_columns = ("id",) if "id" in definition["required"] else sorted(definition["required"])
        seen_keys = set()
        for row_number, item in enumerate(sheet_rows, start=2):
            key = tuple(str(item[column]) for column in key_columns)
            if key in seen_keys:
                raise ValueError(
                    f"{sheet_name} row {row_number}: duplicate key ({', '.join(key)})"
                )
            seen_keys.add(key)
        parsed[sheet_name] = sheet_rows
    return parsed


def _collect_known_ids(rows: Mapping[str, List[Dict[str, Any]]], db: Session) -> Dict[Type[Any], set]:
    ids: Dict[Type[Any], set[str]] = {}
    for model in (AcademicYear, Section, Subject, Teacher, Room, LabBatch, TimeSlot, ElectiveGroup, Department):
        existing = {str(item[0]) for item in db.query(model.id).all()}
        sheet_name = next((name for name, definition in SHEET_DEFINITIONS.items() if definition["model"] is model), "")
        existing.update(str(row["id"]) for row in rows.get(sheet_name, []) if row.get("id") is not None)
        ids[model] = existing
    return ids


def _validate_child_patches(
    child_values: Mapping[str, Mapping[str, Mapping[str, Any]]], ids: Dict[Type[Any], set]
) -> None:
    """child_values holds only the rows _read_rows couldn't attach to a row
    in this same workbook (e.g. a TeacherAvailability row for a teacher that
    already exists in the DB but wasn't re-listed on the Teachers sheet).
    Those still need their parent id, and any secondary reference inside the
    built value (e.g. ElectiveGroupSubjects' subject_id), checked against
    something that actually exists."""
    errors: List[str] = []
    for parent_sheet, per_id in child_values.items():
        model = SHEET_DEFINITIONS[parent_sheet]["model"]
        for parent_id, fields in per_id.items():
            if parent_id not in ids[model]:
                errors.append(f"{parent_sheet}: no existing record '{parent_id}' for the linked child-sheet rows")
    for child_name, (ref_model, label) in CHILD_VALUE_REFERENCES.items():
        definition = CHILD_SHEET_DEFINITIONS[child_name]
        per_id = child_values.get(definition["parent_sheet"], {})
        for parent_id, fields in per_id.items():
            for value in fields.get(definition["parent_field"]) or []:
                if str(value) not in ids[ref_model]:
                    errors.append(f"{child_name}: unknown {label} '{value}' linked to '{parent_id}'")
    if errors:
        raise ValueError("; ".join(errors[:20]) + ("; ..." if len(errors) > 20 else ""))


def _validate_references(rows: Mapping[str, List[Dict[str, Any]]], db: Session, ids: Dict[Type[Any], set]) -> None:
    known_subject_types = {row[0] for row in db.query(SubjectType.id).all()}
    known_subject_types.update(row["type"] for row in rows.get("Subjects", []) if row.get("type"))

    references = {
        "Sections": (("year_id", AcademicYear, "academic year"),),
        "SectionSubjects": (("section_id", Section, "section"), ("subject_id", Subject, "subject")),
        "TeacherSubjects": (("teacher_id", Teacher, "teacher"), ("subject_id", Subject, "subject")),
        "LabBatches": (("section_id", Section, "section"),),
        "Prerequisites": (("subject_id", Subject, "subject"), ("requires_subject_id", Subject, "subject")),
        "ElectiveGroups": (("department_id", Department, "department"),),
    }
    # Nullable FKs: only validated when the cell is actually filled in - a
    # subject or section-subject need not belong to an elective basket, and
    # AcademicYear/Subject/Teacher/Room's department_id is optional too
    # (None = institution-wide / unassigned). But a WRONG id (the id column
    # was filled in with something that doesn't exist) must still fail here
    # with a readable message instead of a raw sqlite3.IntegrityError at
    # insert time - that's what actually happened with a workbook that used
    # department_id "default" against a database that already had a
    # department seeded under a different id.
    optional_references = {
        "Subjects": (
            ("elective_group_id", ElectiveGroup, "elective group"),
            ("department_id", Department, "department"),
        ),
        "SectionSubjects": (("elective_group_id", ElectiveGroup, "elective group"),),
        "AcademicYears": (("department_id", Department, "department"),),
        "Teachers": (("department_id", Department, "department"),),
        "Rooms": (("department_id", Department, "department"),),
    }
    errors: List[str] = []
    for sheet_name, fields in references.items():
        for row_number, row in enumerate(rows.get(sheet_name, []), start=2):
            for field, model, label in fields:
                if str(row[field]) not in ids[model]:
                    errors.append(f"{sheet_name} row {row_number}: unknown {label} '{row[field]}'")
    for sheet_name, fields in optional_references.items():
        for row_number, row in enumerate(rows.get(sheet_name, []), start=2):
            for field, model, label in fields:
                value = row.get(field)
                if value and str(value) not in ids[model]:
                    errors.append(f"{sheet_name} row {row_number}: unknown {label} '{value}'")

    for row_number, row in enumerate(rows.get("Subjects", []), start=2):
        if row.get("type") and row["type"] not in known_subject_types:
            errors.append(f"Subjects row {row_number}: unknown subject type '{row['type']}' (add it under Subject Types first)")

    for row_number, row in enumerate(rows.get("Rooms", []), start=2):
        if row.get("type") and row["type"] not in VALID_ROOM_TYPES:
            errors.append(f"Rooms row {row_number}: type must be one of {sorted(VALID_ROOM_TYPES)}, got '{row['type']}'")

    for row_number, row in enumerate(rows.get("ElectiveGroups", []), start=2):
        for subject_id in row.get("offered_subject_ids") or []:
            if str(subject_id) not in ids[Subject]:
                errors.append(f"ElectiveGroups row {row_number}: unknown subject '{subject_id}' in offered_subject_ids")

    if errors:
        raise ValueError("; ".join(errors[:20]) + ("; ..." if len(errors) > 20 else ""))


def _validate_schemas(rows: Mapping[str, List[Dict[str, Any]]]) -> None:
    """Runs every parsed row through the same Pydantic schema the single-row
    API endpoints use. Before this, the importer built SQLAlchemy models
    directly from the parsed cells - a workbook with an invalid
    delivery_mode, a day outside Mon-Sun, or a lecture/tutorial/practical
    split that doesn't add up to weekly_hours imported "successfully" and
    only surfaced as an empty or wrong timetable at generation time."""
    errors: List[str] = []
    for sheet_name, definition in SHEET_DEFINITIONS.items():
        schema = definition.get("schema")
        if schema is None:
            continue
        for row_number, row in enumerate(rows.get(sheet_name, []), start=2):
            try:
                schema(**row)
            except ValidationError as exc:
                for err in exc.errors():
                    field = ".".join(str(part) for part in err["loc"]) or sheet_name
                    errors.append(f"{sheet_name} row {row_number}: {field}: {err['msg']}")
    if errors:
        raise ValueError("; ".join(errors[:20]) + ("; ..." if len(errors) > 20 else ""))


def _upsert_entities(
    rows: Mapping[str, List[Dict[str, Any]]],
    db: Session,
    child_patches: Mapping[str, Mapping[str, Mapping[str, Any]]] = {},
) -> Dict[str, int]:
    counts = {"created": 0, "updated": 0, "relations_created": 0}
    entity_sheets = [name for name in SHEET_DEFINITIONS if name not in RELATION_KEYS]
    for sheet_name in entity_sheets:
        definition = SHEET_DEFINITIONS[sheet_name]
        model = definition["model"]
        for data in rows.get(sheet_name, []):
            record = db.get(model, str(data["id"]))
            if record is None:
                record = model(**data)
                db.add(record)
                counts["created"] += 1
            else:
                for key, value in data.items():
                    setattr(record, key, value)
                counts["updated"] += 1
    db.flush()

    for sheet_name, (model, key_fields) in RELATION_KEYS.items():
        for data in rows.get(sheet_name, []):
            filters = {field: data[field] for field in key_fields}
            record = db.query(model).filter_by(**filters).first()
            if record is None:
                db.add(model(**data))
                counts["relations_created"] += 1
            else:
                for key, value in data.items():
                    setattr(record, key, value)

    # Child-sheet rows whose parent wasn't re-listed on this workbook (the
    # parent already exists from an earlier import) - patch it directly
    # instead of round-tripping it through the Create schema, which would
    # fail on required fields (e.g. Teacher.name) this row never carried.
    for parent_sheet, per_id in child_patches.items():
        model = SHEET_DEFINITIONS[parent_sheet]["model"]
        for parent_id, fields in per_id.items():
            record = db.get(model, parent_id)
            if record is None:
                continue
            for key, value in fields.items():
                setattr(record, key, value)
            counts["updated"] += 1
    return counts


def _parse_and_validate_workbook(content: bytes, db: Session) -> tuple[Dict[str, List[Dict[str, Any]]], Dict[str, Dict[str, Dict[str, Any]]]]:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    child_values = _read_child_sheets(workbook)
    rows = _read_rows(workbook, child_values)
    if not rows and not any(child_values.values()):
        raise ValueError("The workbook does not contain any supported data sheets")
    ids = _collect_known_ids(rows, db)
    _validate_child_patches(child_values, ids)
    _validate_references(rows, db, ids)
    _validate_schemas(rows)
    return rows, child_values


@router.post("/excel/preview")
async def preview_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin: User = Depends(require_admin),
):
    """Parse and validate an upload without changing the database."""
    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload an .xlsx or .xlsm workbook")
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded workbook is empty")
    try:
        rows, child_values = _parse_and_validate_workbook(content, db)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Workbook preview failed: {exc}") from exc
    linked_updates = sum(len(per_id) for per_id in child_values.values())
    return {
        "message": "Workbook validated. No data has been imported.",
        "sheets": {name: len(sheet_rows) for name, sheet_rows in rows.items()},
        "total_rows": sum(len(sheet_rows) for sheet_rows in rows.values()),
        "linked_updates": linked_updates,
    }


@router.post("/csv")
async def import_csv_bundle(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin: User = Depends(require_admin),
):
    """Import either one named CSV sheet or a ZIP bundle of named CSV sheets."""
    filename = (file.filename or "").lower()
    if not filename.endswith((".csv", ".zip")):
        raise HTTPException(status_code=400, detail="Upload a .csv file or a .zip bundle of CSV sheets")
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded CSV file is empty")
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        sources = {}
        if filename.endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                for member in archive.namelist():
                    if member.lower().endswith(".csv") and not member.endswith("/"):
                        sources[member.rsplit("/", 1)[-1]] = archive.read(member).decode("utf-8-sig")
        else:
            sources[file.filename or "data.csv"] = content.decode("utf-8-sig")
        supported = {f"{name.lower()}.csv": name for name in list(SHEET_DEFINITIONS) + list(CHILD_SHEET_DEFINITIONS)}
        for source_name, source_text in sources.items():
            sheet_name = supported.get(source_name.lower())
            if not sheet_name:
                continue
            sheet = workbook.create_sheet(sheet_name)
            for row in csv.reader(io.StringIO(source_text)):
                sheet.append(row)
        if not workbook.sheetnames:
            raise ValueError("CSV names must match supported sheets, for example Subjects.csv or TeacherAvailability.csv")
        child_values = _read_child_sheets(workbook)
        rows = _read_rows(workbook, child_values)
        ids = _collect_known_ids(rows, db)
        _validate_child_patches(child_values, ids)
        _validate_references(rows, db, ids)
        _validate_schemas(rows)
        counts = _upsert_entities(rows, db, child_values)
        db.commit()
    except (ValueError, UnicodeDecodeError, zipfile.BadZipFile) as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail=f"CSV import failed: {exc}") from exc
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail=f"CSV import failed: {exc}") from exc
    return {
        "message": "CSV data imported successfully",
        "sheets": {name: len(sheet_rows) for name, sheet_rows in rows.items()},
        **counts,
    }


@router.post("/excel")
async def import_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin: User = Depends(require_admin),
):
    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload an .xlsx or .xlsm workbook")
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded workbook is empty")
    try:
        rows, child_values = _parse_and_validate_workbook(content, db)
        counts = _upsert_entities(rows, db, child_values)
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail=f"Workbook import failed: {exc}") from exc

    return {
        "message": "Workbook imported successfully",
        "sheets": {name: len(sheet_rows) for name, sheet_rows in rows.items()},
        **counts,
    }


# One example row per child sheet, shown in the template so admins see the
# table shape instead of guessing it from a column-header list alone.
CHILD_SHEET_EXAMPLES: Dict[str, List[Any]] = {
    "TeacherAvailability": ["T1", "Mon", "09:00", "13:00"],
    "TeacherPreferredSlots": ["T1", "Tue", "10:00", "12:00"],
    "RoomAvailability": ["R1", "Mon", "08:00", "18:00"],
    "RoomEquipment": ["R1", "projector"],
    "RoomSharedDepartments": ["R1", "CSE"],
    "SubjectEquipment": ["S1", "projector"],
    "ElectiveGroupSubjects": ["EG1", "S1"],
    "AcademicYearLunchWindows": ["Y1", "Mon", "12:30", "13:15"],
}


@router.get("/excel-template")
def download_excel_template(_admin: User = Depends(require_admin)):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["Bulk import workbook", "Fill one or more supported sheets and upload the .xlsx file."])
    instructions.append([
        "Linking sheets",
        "Fields that hold several values (a teacher's free periods, a room's equipment, an elective "
        "group's subjects, ...) are their own sheet - one row per value, linked back by id. "
        "For example, add one row per free window to TeacherAvailability instead of one JSON cell on Teachers.",
    ])
    instructions.append(["Time fields", "Use HH:MM values, for example 09:00."])
    for sheet_name, definition in SHEET_DEFINITIONS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(definition["columns"])
        sheet.freeze_panes = "A2"
    for sheet_name, definition in CHILD_SHEET_DEFINITIONS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(definition["columns"])
        example = CHILD_SHEET_EXAMPLES.get(sheet_name)
        if example:
            sheet.append(example)
        sheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=timetable-import-template.xlsx"},
    )
