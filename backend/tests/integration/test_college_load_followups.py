"""The department-load workbook is the main import format, so beyond reading
rows it must get four things right: elective baskets, re-importing over an
admin's corrections, syncing allotments to a newer revision, and giving labs
a batch mode their teachers can actually staff. Plus the blank template."""
import io

from openpyxl import Workbook

ADMIN = {"email": "admin@test.edu", "password": "TestPassword123!", "full_name": "Admin", "role": "SUPERADMIN"}
HEADER = ["S. No.", "Class/ Year", "Subject Code", "Subject Nomenclature", None, "Theory", "Tutorial", "Lab"]


def _headers(client):
    assert client.post("/api/auth/bootstrap-admin", json=ADMIN).status_code == 200
    r = client.post("/api/auth/login", json={"email": ADMIN["email"], "password": ADMIN["password"]})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _workbook(rows, title="Subject_facult_19Aug26", date_text="19/08/2026") -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    faculty = workbook.create_sheet("faculty_list")
    faculty.append(["Department of Computer Engineering"])
    faculty.append(["S.No.", "Name", None, "Contact No."])
    for number, (name, initials) in enumerate(
        [("Dr. Urjita Thakar", "UT"), ("Ms. Swati Mishra", "SM"), ("Ms. Himani Mishra", "HM")], start=1
    ):
        faculty.append([number, name, initials])
    sheet = workbook.create_sheet(title)
    sheet.append(["Department of Computer Engineering"])
    sheet.append([f"Date: {date_text}"])
    sheet.append([])
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _import(client, headers, content, path="/api/import/excel"):
    r = client.post(path, headers=headers, files={"file": ("load.xlsx", content, "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


def _subjects(client, headers):
    return {s["id"]: s for s in client.get("/api/subjects/", headers=headers).json()}


def test_elective_basket_holds_only_the_main_part_of_each_elective(client):
    headers = _headers(client)
    _import(client, headers, _workbook([
        [1, "III B", "CO34501", "Data Base Management Systems", "3L", "UT(3)"],
        [None, None, "CO3____", "Elective-I: Data Science", "3L+P", "UT(3)", None, "HM(2), SM(2)"],
        [None, None, "CO3____", "Elective-I: Artificial Intelligence", "3L", "SM(3)"],
    ]))
    groups = client.get("/api/elective-groups/", headers=headers).json()
    assert len(groups) == 1
    by_name = {s["name"]: s for s in _subjects(client, headers).values()}
    theory = by_name["Elective-I: Data Science (Theory)"]
    lab = by_name["Elective-I: Data Science Lab"]
    ai = by_name["Elective-I: Artificial Intelligence"]
    # Alternatives share the basket; the lab the same students attend does not.
    assert set(groups[0]["offered_subject_ids"]) == {theory["id"], ai["id"]}
    assert theory["elective_group_id"] == groups[0]["id"]
    assert lab["elective_group_id"] is None
    assert lab["linked_group_id"] == theory["id"]  # still tied to its theory


def test_reimport_keeps_what_the_admin_corrected(client):
    headers = _headers(client)
    content = _workbook([
        [1, "II A", "CO24402", "Computer Architecture Lab", "1P", None, None, "UT(6), HM(6)"],
        [2, "PG I", "CO71016", "Programming Systems", "3L", "ZZ(3)"],
    ])
    _import(client, headers, content)
    assert len(client.get("/api/constraints/lab-batches/SEC_II_A", headers=headers).json()) == 3

    # The admin fixes the placeholders the importer had to invent.
    section = client.get("/api/sections/SEC_II_A", headers=headers).json()
    assert client.put("/api/sections/SEC_II_A", headers=headers, json={**section, "strength": 75}).status_code == 200
    teacher = client.get("/api/teachers/T_ZZ", headers=headers).json()
    assert client.put("/api/teachers/T_ZZ", headers=headers, json={**teacher, "name": "Zed Zee"}).status_code == 200
    first_batch = client.get("/api/constraints/lab-batches/SEC_II_A", headers=headers).json()[0]["id"]
    assert client.delete(f"/api/constraints/lab-batches/{first_batch}", headers=headers).status_code == 200

    _import(client, headers, content)

    assert client.get("/api/sections/SEC_II_A", headers=headers).json()["strength"] == 75
    assert client.get("/api/teachers/T_ZZ", headers=headers).json()["name"] == "Zed Zee"
    assert len(client.get("/api/constraints/lab-batches/SEC_II_A", headers=headers).json()) == 2


def test_reimport_keeps_an_elective_basket_the_admin_made_non_parallel(client):
    headers = _headers(client)
    content = _workbook([
        [1, "III B", "CO3____", "Elective-I: Data Science", "3L", "UT(3)"],
        [None, None, "CO3____", "Elective-I: Artificial Intelligence", "3L", "SM(3)"],
    ])
    _import(client, headers, content)
    group = client.get("/api/elective-groups/", headers=headers).json()[0]
    assert group["must_be_parallel"] is True  # a new basket starts parallel
    assert client.put(f"/api/elective-groups/{group['id']}", headers=headers,
                      json={**group, "must_be_parallel": False}).status_code == 200

    _import(client, headers, content)

    assert client.get("/api/elective-groups/", headers=headers).json()[0]["must_be_parallel"] is False


def test_reimport_of_a_newer_revision_syncs_allotments_and_leaves_manual_data(client):
    headers = _headers(client)
    first = _workbook([
        [1, "II A", "CO24019", "Computer Architecture", "3L", "HM(3)"],
        [None, None, "CO24402", "Computer Architecture Lab", "1P", None, None, "HM(6), SM(6)"],
    ], title="Subject_facult_15Jun26", date_text="15/06/2026")
    newer = _workbook([
        [1, "II A", "CO24019", "Computer Architecture", "3L", "UT(3)"],
        [None, None, "CO24402", "Computer Architecture Lab", "1P", None, None, "SM(6), UT(6)"],
    ])
    _import(client, headers, first)

    # Things the workbook knows nothing about.
    assert client.post("/api/rooms/", headers=headers, json={
        "id": "R1", "name": "Room 1", "type": "lecture", "capacity": 90, "department_id": "default"}).status_code == 200
    assert client.post("/api/constraints/time-slots", headers=headers, json={
        "id": "S1", "day": "Mon", "period_index": 1, "start_time": "09:00", "end_time": "10:00",
        "department_id": "default"}).status_code == 200
    assert client.post("/api/subjects/", headers=headers, json={
        "id": "HAND", "name": "Added by hand", "type": "theory", "weekly_hours": 1,
        "department_id": "default"}).status_code == 200
    assert client.post("/api/constraints/section-subjects", headers=headers,
                       json={"section_id": "SEC_II_A", "subject_id": "HAND"}).status_code == 200

    preview = client.post("/api/import/excel/preview", headers=headers,
                          files={"file": ("load.xlsx", newer, "application/octet-stream")}).json()
    assert any("no longer in the workbook" in w for w in preview["warnings"])  # told before confirming
    _import(client, headers, newer)

    def qualified(teacher_id):
        return set(client.get(f"/api/teachers/{teacher_id}/subjects", headers=headers).json())

    assert "CO24019" in qualified("T_UT") and "CO24019" not in qualified("T_HM")
    assert "CO24402" not in qualified("T_HM")
    assert "CO24402" in qualified("T_UT") and "CO24402" in qualified("T_SM")
    assert [r["id"] for r in client.get("/api/rooms/", headers=headers).json()] == ["R1"]
    assert [s["id"] for s in client.get("/api/constraints/time-slots", headers=headers).json()] == ["S1"]
    assigned = {r["subject_id"] for r in client.get("/api/constraints/section-subjects/SEC_II_A", headers=headers).json()}
    assert "HAND" in assigned  # not in the workbook, so not the import's to remove


def test_labs_get_a_batch_mode_the_teachers_can_actually_staff(client):
    headers = _headers(client)
    # II A: three 2-period labs, one teacher each, three batches -> rotate.
    # II B: only two labs for three batches -> take the batches one after another.
    content = _workbook([
        [1, "II A", "CO24401", "OOP Lab", "1P", None, None, "UT(6)"],
        [None, None, "CO24402", "Architecture Lab", "1P", None, None, "HM(6)"],
        [None, None, "CO24404", "Design Thinking Lab", "1P", None, None, "SM(6)"],
        [2, "II B", "CO24801", "DS Lab", "1P", None, None, "UT(6)"],
        [None, None, "CO24802", "Agile Lab", "1P", None, None, "HM(6)"],
    ])
    _import(client, headers, content)
    rotations = client.get("/api/lab-rotations/", headers=headers).json()
    assert [r["section_id"] for r in rotations] == ["SEC_II_A"]
    assert rotations[0]["source"] == "import"
    assert set(rotations[0]["subject_ids"]) == {"CO24401", "CO24402", "CO24404"}
    subjects = _subjects(client, headers)
    assert subjects["CO24801"]["batch_scheduling_mode"] == "sequential"
    assert subjects["CO24802"]["batch_scheduling_mode"] == "sequential"
    assert subjects["CO24401"]["batch_scheduling_mode"] == "independent"  # it rotates, so no change

    # Running it again changes nothing and doesn't pile up duplicates.
    _import(client, headers, content)
    assert len(client.get("/api/lab-rotations/", headers=headers).json()) == 1


def test_blank_template_is_in_the_same_format_and_reads_back(client):
    headers = _headers(client)
    r = client.get("/api/import/college-template", headers=headers)
    assert r.status_code == 200, r.text
    assert "department-load-template.xlsx" in r.headers["content-disposition"]
    preview = client.post("/api/import/excel/preview", headers=headers,
                          files={"file": ("department-load-template.xlsx", r.content, "application/octet-stream")})
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["format"] == "college_load"
    assert body["sheets"]["Teachers"] == 3 and body["sheets"]["Sections"] == 3
    # the Load Calculation's "No. of Batches" gave the lab its 3 batches
    assert body["sheets"]["LabBatches"] == 3
