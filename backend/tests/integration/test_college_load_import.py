"""The department's own teaching-load workbook (faculty list + several dated
revisions of the subject/faculty allotment) is converted to the bulk-import
template on upload. These tests use a small workbook laid out like the
real one: merged-style title rows, a class label only on the first row of
each block, re-registration ("RR") rows under the last block, and
teachers referred to by initials with hours."""
import io

from openpyxl import Workbook

ADMIN = {"email": "admin@test.edu", "password": "TestPassword123!", "full_name": "Admin", "role": "SUPERADMIN"}
HEADER = ["S. No.", "Class/ Year", "Subject Code", "Subject Nomenclature", None, "Theory", "Tutorial", "Lab"]


def _headers(client):
    assert client.post("/api/auth/bootstrap-admin", json=ADMIN).status_code == 200
    r = client.post("/api/auth/login", json={"email": ADMIN["email"], "password": ADMIN["password"]})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _allotment_sheet(workbook, title, date_text, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(["Department of Computer Engineering"])
    sheet.append([f"Date: {date_text}"])
    sheet.append([])
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)


def _college_workbook() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    faculty = workbook.create_sheet("faculty_list")
    faculty.append(["Department of Computer Engineering"])
    faculty.append(["S.No.", "Name", None, "Contact No."])
    faculty.append([1, "Dr. Urjita Thakar", "UT"])
    faculty.append([2, "Ms. Swati Mishra", "SM"])
    faculty.append([3, "Ms. Himani Mishra", "HM"])
    faculty.append(["Lab Technicians"])
    faculty.append([1, "Mr. Sudheer Anwekar", "SA'"])

    # Older revision: Computer Architecture was with UT then.
    _allotment_sheet(workbook, "subject_facult_15June26", "15/06/2026", [
        [1, "II A", "CO24019", "Computer Architecture", "3L", "UT(3)"],
    ])
    _allotment_sheet(workbook, "Subject_facult_19Aug26", "19/08/2026", [
        [1, "II A", "CO24019", "Computer Architecture", "3L", "HM(3)"],
        [None, None, "CO24017", "Object Oriented Programming", "3L+1Tut.", "UT(3)", "SwM(1)"],
        [None, None, "CO24402", "Computer Architecture Lab", "1P", None, None, "HM(6), SwM(6)"],
        [None, None, "CO24498", "Mini Project", "2P"],
        [None, None, "CO24002 RR", "Theory of Computation (3L+1Tut.)", None, "UT(1)"],
        [2, "PG I", "CO71016", "Programming Systems", "3L", "ZZ(3)"],
        [3, None, None, "IT10007", None, "UT(2)"],
    ])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _upload(client, headers, path):
    return client.post(
        path, headers=headers,
        files={"file": ("01_Load_july2026.xlsx", _college_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_college_load_workbook_is_previewed_from_latest_revision(client):
    headers = _headers(client)
    r = _upload(client, headers, "/api/import/excel/preview")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["format"] == "college_load"
    assert body["source_sheet"] == "Subject_facult_19Aug26"
    warnings = " ".join(body["warnings"])
    assert "subject_facult_15June26" in warnings
    assert "'ZZ'" in warnings  # unknown initials are flagged, not dropped
    assert "IT10007" in warnings  # rows with no class are reported as skipped
    assert body["sheets"]["Sections"] == 3  # II A, II RR, PG I


def test_college_load_workbook_imports_sections_subjects_and_teachers(client):
    headers = _headers(client)
    r = _upload(client, headers, "/api/import/excel")
    assert r.status_code == 200, r.text

    sections = {s["id"]: s for s in client.get("/api/sections/", headers=headers).json()}
    assert set(sections) == {"SEC_II_A", "SEC_II_RR", "SEC_PG_I"}

    subjects = {s["id"]: s for s in client.get("/api/subjects/", headers=headers).json()}
    # L+T becomes linked theory + tutorial rows.
    assert subjects["CO24017"]["type"] == "theory" and subjects["CO24017"]["weekly_hours"] == 3
    assert subjects["CO24017_TUT"]["type"] == "tutorial" and subjects["CO24017_TUT"]["weekly_hours"] == 1
    assert subjects["CO24017_TUT"]["linked_group_id"] == "CO24017"
    # 1P = one 2-period lab per batch.
    assert subjects["CO24402"]["type"] == "lab" and subjects["CO24402"]["weekly_hours"] == 2
    # RR offering is separate, at the hours actually allotted.
    assert subjects["CO24002_RR"]["weekly_hours"] == 1
    # Nobody allotted yet -> kept, but not scheduled.
    assert subjects["CO24498"]["delivery_mode"] == "SELF_STUDY"

    teachers = {t["id"]: t for t in client.get("/api/teachers/", headers=headers).json()}
    assert "T_SA" not in teachers  # lab technicians aren't teaching faculty
    assert teachers["T_ZZ"]["name"] == "ZZ"

    def qualified(teacher_id):
        return set(client.get(f"/api/teachers/{teacher_id}/subjects", headers=headers).json())

    # Latest revision wins: Computer Architecture is Himani's, not Urjita's.
    assert "CO24019" in qualified("T_HM")
    assert "CO24019" not in qualified("T_UT")
    # "SwM" in the allotment is Swati Mishra ("SM" in the faculty list).
    assert {"CO24017_TUT", "CO24402"} <= qualified("T_SM")

    # Lab batches aren't in the workbook: A/B sections get the usual 3.
    batches = client.get("/api/constraints/lab-batches/SEC_II_A", headers=headers)
    assert batches.status_code == 200, batches.text
    assert len(batches.json()) == 3


def test_college_load_workbook_goes_into_the_selected_department(client):
    """The workbook has no department id of its own - the years must land
    in the department the admin has selected, or the Years page (which is
    filtered by that department) shows none of them."""
    headers = _headers(client)
    client.post("/api/departments/", json={"id": "mine", "institution_id": "inst1", "name": "My Department"},
                headers=headers)
    r = client.post(
        "/api/import/excel", headers=headers, data={"department_id": "mine"},
        files={"file": ("load.xlsx", _college_workbook(), "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    assert r.json()["department_name"] == "My Department"
    years = client.get("/api/years/", headers=headers).json()
    assert years and {y["department_id"] for y in years} == {"mine"}


def test_template_workbook_is_not_treated_as_college_format(client):
    headers = _headers(client)
    workbook = Workbook()
    workbook.active.title = "Teachers"
    workbook.active.append(["id", "name"])
    workbook.active.append(["T1", "Teacher One"])
    output = io.BytesIO()
    workbook.save(output)
    r = client.post("/api/import/excel/preview", headers=headers,
                    files={"file": ("t.xlsx", output.getvalue(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert r.json()["format"] == "template"
